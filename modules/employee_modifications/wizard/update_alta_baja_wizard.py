# -*- coding: utf-8 -*-

import base64
from datetime import date, datetime, timedelta
from io import BytesIO

import xlrd
from openpyxl import load_workbook

from odoo import models, fields, _
from odoo.exceptions import UserError, AccessError
from markupsafe import escape


class UpdateAltaBajaWizard(models.TransientModel):
    _name = 'update.alta.baja.wizard'
    _description = 'Actualizar fechas de alta y baja desde Excel'

    archivo_excel = fields.Binary(string='Archivo Excel', required=True)
    nombre_archivo = fields.Char(string='Nombre del Archivo')
    resultado = fields.Html(string='Resultado', readonly=True)
    state = fields.Selection([
        ('seleccionar', 'Seleccionar Archivo'),
        ('resultado', 'Resultado'),
    ], default='seleccionar')
    actualizados = fields.Integer(readonly=True)
    omitidos = fields.Integer(readonly=True)

    def _normalize_text(self, value):
        if value is None:
            return ''
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    def _canonical_number(self, value):
        """Ignora ceros a la izquierda en números (0008 == 8)."""
        value = (value or '').strip()
        if value.isdigit():
            return value.lstrip('0') or '0'
        return value.upper()

    def _find_employees(self, Employee, numero):
        target = self._canonical_number(numero)
        suffix = numero.lstrip('0') if numero.isdigit() else numero
        candidates = Employee.search([
            ('biometric_id', '=ilike', f'%{suffix}'),
            ('company_id', '=', self.env.company.id),
        ])
        return candidates.filtered(
            lambda e: self._canonical_number(e.biometric_id) == target
        )

    def _parse_date(self, value, datemode=0):
        """Devuelve (fecha|None, error|None). Vacío => (None, None)."""
        if value is None or (isinstance(value, str) and not value.strip()):
            return None, None
        if isinstance(value, datetime):
            return value.date(), None
        if isinstance(value, date):
            return value, None
        if isinstance(value, (int, float)):
            try:
                if datemode is None:
                    return date(1899, 12, 30) + timedelta(days=float(value)), None
                return xlrd.xldate.xldate_as_datetime(value, datemode).date(), None
            except Exception:
                return None, _('Fecha inválida: %s') % value
        text = str(value).strip()
        for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y', '%Y/%m/%d'):
            try:
                return datetime.strptime(text, fmt).date(), None
            except ValueError:
                continue
        return None, _('Fecha inválida: %s') % text

    def _read_rows(self, content, file_name):
        """Devuelve (filas, datemode). datemode None => openpyxl."""
        name = (file_name or '').lower()

        def read_xlsx():
            wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
            rows = [r for r in wb.worksheets[0].iter_rows(min_row=2, max_col=3, values_only=True)]
            return rows, None

        def read_xls():
            wb = xlrd.open_workbook(file_contents=content)
            sheet = wb.sheet_by_index(0)
            rows = [
                tuple(sheet.cell_value(r, c) if sheet.ncols > c else '' for c in range(3))
                for r in range(1, sheet.nrows)
            ]
            return rows, wb.datemode

        if name.endswith('.xlsx'):
            return read_xlsx()
        if name.endswith('.xls'):
            return read_xls()
        try:
            return read_xlsx()
        except Exception:
            return read_xls()

    def action_actualizar(self):
        self.ensure_one()
        if not self.env.user.has_group('base.group_system'):
            raise AccessError(_('Solo el administrador puede utilizar esta herramienta.'))
        if not self.archivo_excel:
            raise UserError(_('Por favor seleccione un archivo Excel.'))

        try:
            rows, datemode = self._read_rows(base64.b64decode(self.archivo_excel), self.nombre_archivo)
        except Exception as e:
            raise UserError(_('No se pudo leer el archivo Excel: %s') % e)

        Employee = self.env['hr.employee'].with_context(active_test=False)
        actualizados = 0
        omitidos = 0
        html = ("<table class='table table-striped'><thead><tr><th>Fila</th>"
                "<th>No. Empleado</th><th>Estado</th><th>Detalles</th></tr></thead><tbody>")

        for idx, row in enumerate(rows, start=2):
            numero = self._normalize_text(row[0] if len(row) > 0 else '')
            if not numero:
                continue
            estado, detalle = 'Omitido', ''
            fecha_alta, err_alta = self._parse_date(row[1] if len(row) > 1 else None, datemode)
            fecha_baja, err_baja = self._parse_date(row[2] if len(row) > 2 else None, datemode)

            if err_alta or err_baja:
                detalle = err_alta or err_baja
            elif not fecha_alta and not fecha_baja:
                detalle = _('Sin fechas para actualizar')
            else:
                employees = self._find_employees(Employee, numero)
                if not employees:
                    detalle = _('Empleado no encontrado')
                elif len(employees) > 1:
                    detalle = _('Número de empleado duplicado')
                elif employees.active:
                    detalle = _('El empleado está activo; solo se actualizan inactivos')
                else:
                    vals = {}
                    if fecha_alta:
                        vals['fecha_ingreso_manual'] = fecha_alta
                    if fecha_baja:
                        vals['departure_date'] = fecha_baja
                    employees.with_context(skip_fecha_ingreso_sync=True).write(vals)
                    actualizados += 1
                    estado = 'Actualizado'
                    detalle = ', '.join(f'{k}: {v}' for k, v in vals.items())
            if estado == 'Omitido':
                omitidos += 1
            html += (f"<tr><td>{idx}</td><td>{escape(numero)}</td>"
                     f"<td>{estado}</td><td>{escape(detalle)}</td></tr>")
        html += "</tbody></table>"

        self.write({
            'resultado': html,
            'actualizados': actualizados,
            'omitidos': omitidos,
            'state': 'resultado',
        })
        return self._reopen()

    def action_volver(self):
        self.write({'state': 'seleccionar', 'archivo_excel': False, 'resultado': False})
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
